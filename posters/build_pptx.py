"""Builds an editable one-slide PPTX from extract.js output: textless background image plus
one text box per text block (RTL Arabic, original fonts/colours/sizes). Import it into Canva.
Usage: python3 build_pptx.py <name>   (reads out/<name>-bg.png and out/<name>-text.json)
"""
import json, sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

PX = 9525  # EMU per CSS pixel

name = sys.argv[1]
data = json.load(open(f'out/{name}-text.json'))
W, H = 1500, data['height']

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.shapes.add_picture(f'out/{name}-bg.png', 0, 0, Emu(W * PX), Emu(H * PX))

def set_fonts(run, font):
    rPr = run._r.get_or_add_rPr()
    rPr.set('lang', 'ar-SA')
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set('typeface', font)

for b in data['blocks']:
    x1, x2 = b['x'], b['x2']
    if b['lines'] > 1 and (b['align'] == 'center' or b['hasBefore']):
        x1, x2 = b['cx1'], b['cx2']            # wrapped text: use the full column width
    elif b['lines'] > 1:
        x1, x2 = b['cx1'], b['firstRight'] + 4  # keep clear of a drawn bullet on the right
    elif b['align'] == 'center':
        pad = 0.18 * (x2 - x1) + 30
        x1, x2 = x1 - pad, x2 + pad
    else:
        x1 = b['cx1']                          # room to grow leftwards without re-wrapping
        x2 = b['cx2'] if b['hasBefore'] else x2 + 4
    lh = b['lh']
    size = max(r.get('size', 0) for r in b['runs'])
    top = b['y'] - (lh - (b['y2'] - b['y']) / b['lines']) / 2 - 0.15 * size
    box = slide.shapes.add_textbox(Emu(int(x1 * PX)), Emu(int(top * PX)),
                                   Emu(int((x2 - x1) * PX)), Emu(int((lh * b['lines'] + 6) * PX)))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    # split runs into paragraphs on explicit line breaks
    paras, cur = [], []
    for r in b['runs']:
        if r['text'] == '\n':
            paras.append(cur); cur = []
        else:
            cur.append(r)
    paras.append(cur)

    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER if b['align'] == 'center' else PP_ALIGN.RIGHT
        p._p.get_or_add_pPr().set('rtl', '1')
        p.line_spacing = Pt(lh * 0.75)
        if runs:
            runs[0] = dict(runs[0], text=runs[0]['text'].lstrip())
            runs[-1] = dict(runs[-1], text=runs[-1]['text'].rstrip())
        for r in runs:
            if not r['text']:
                continue
            run = p.add_run()
            run.text = r['text']
            run.font.size = Pt(r['size'] * 0.75)
            run.font.bold = r['weight'] >= 600
            run.font.color.rgb = RGBColor.from_string(r['color'])
            set_fonts(run, r['font'])

prs.save(f'out/{name}.pptx')
print(f'out/{name}.pptx', len(data['blocks']), 'text boxes')
